"""Build an Excel workbook from ``SystemVersionContent``."""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.system.content import SystemVersionContent
from app.domain.system.entities import System
from app.domain.system.entities import SystemVersion
from app.infrastructure.excel import schema as s


def _write_header(ws: Worksheet, columns: tuple[str, ...]) -> None:
    for col_idx, name in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=name)


def _write_rows(
    ws: Worksheet,
    columns: tuple[str, ...],
    rows: list[dict[str, Any]],
) -> None:
    _write_header(ws, columns)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, key in enumerate(columns, start=1):
            value = row.get(key)
            if value is None:
                continue
            if isinstance(value, bool):
                ws.cell(row=row_idx, column=col_idx, value=str(value))
            else:
                ws.cell(row=row_idx, column=col_idx, value=value)


def _tag_by_id(content: SystemVersionContent) -> dict[UUID, str]:
    return {eq.id: eq.tag for eq in content.equipment}


def _fm_name_by_id(content: SystemVersionContent) -> dict[UUID, str]:
    return {fm.id: fm.name for fm in content.failure_modes}


def _maint_name_by_id(content: SystemVersionContent) -> dict[UUID, str]:
    return {task.id: task.name for task in content.maintenance_tasks}


def _comp_name_by_id(content: SystemVersionContent) -> dict[UUID, str]:
    return {comp.id: comp.name for comp in content.components}


def export_version_workbook(
    system: System,
    version: SystemVersion,
    content: SystemVersionContent,
    *,
    dataset_label: str | None = None,
) -> bytes:
    """Serialize version content to ``.xlsx`` bytes."""
    tags = _tag_by_id(content)
    fm_names = _fm_name_by_id(content)
    maint_names = _maint_name_by_id(content)
    comp_names = _comp_name_by_id(content)

    wb = Workbook()
    # Remove the default sheet; recreate named sheets in order.
    default = wb.active
    if default is not None:
        wb.remove(default)

    _write_rows(
        wb.create_sheet(s.SHEET_SYSTEMS),
        s.SYSTEM_COLUMNS,
        [
            {
                "name": system.name,
                "description": system.description,
                "dataset_label": dataset_label
                or f"version-{version.version_number}",
                "for_software_testing": "true",
            }
        ],
    )

    equip_rows: list[dict[str, Any]] = []
    for eq in content.equipment:
        parent_tag = tags.get(eq.parent_id) if eq.parent_id else None
        equip_rows.append(
            {
                "tag": eq.tag,
                "name": eq.name,
                "description": eq.description,
                "parent_tag": parent_tag,
                "category": eq.category,
                "equipment_class": eq.equipment_class,
                "equipment_type": eq.equipment_type,
                "location": eq.location,
                "quantity": eq.quantity,
                "criticality": str(eq.criticality),
                "operating_mode": str(eq.operating_mode),
                "standby_mode": str(eq.standby_mode),
                "is_repairable": eq.is_repairable,
            }
        )
    _write_rows(
        wb.create_sheet(s.SHEET_EQUIPMENT),
        s.EQUIPMENT_COLUMNS,
        equip_rows,
    )

    comp_rows = [
        {
            "equipment_tag": tags.get(comp.equipment_id, ""),
            "name": comp.name,
            "description": comp.description,
            "quantity": comp.quantity,
        }
        for comp in content.components
    ]
    _write_rows(
        wb.create_sheet(s.SHEET_COMPONENTS),
        s.COMPONENT_COLUMNS,
        comp_rows,
    )

    conn_rows = [
        {
            "from_tag": tags.get(conn.source_id, ""),
            "to_tag": tags.get(conn.target_id, ""),
            "connection_type": str(conn.connection_type),
            "description": conn.description,
        }
        for conn in content.connections
    ]
    _write_rows(
        wb.create_sheet(s.SHEET_CONNECTIONS),
        s.CONNECTION_COLUMNS,
        conn_rows,
    )

    fm_rows: list[dict[str, Any]] = []
    pf_rows: list[dict[str, Any]] = []
    for fm in content.failure_modes:
        eq_tag = tags.get(fm.equipment_id, "")
        comp_name = (
            comp_names.get(fm.component_id) if fm.component_id else None
        )
        fm_rows.append(
            {
                "equipment_tag": eq_tag,
                "name": fm.name,
                "description": fm.description,
                "is_detectable": fm.is_detectable,
                "component_name": comp_name,
            }
        )
        if fm.pf_interval is not None:
            pf_rows.append(
                {
                    "equipment_tag": eq_tag,
                    "failure_mode_name": fm.name,
                    "value": fm.pf_interval.value,
                    "unit": str(fm.pf_interval.unit),
                }
            )
    _write_rows(
        wb.create_sheet(s.SHEET_FAILURE_MODES),
        s.FAILURE_MODE_COLUMNS,
        fm_rows,
    )
    _write_rows(
        wb.create_sheet(s.SHEET_PF_INTERVALS),
        s.PF_INTERVAL_COLUMNS,
        pf_rows,
    )

    dist_by_fm = {
        dist.failure_mode_id: dist for dist in content.failure_distributions
    }
    dist_rows: list[dict[str, Any]] = []
    for fm in content.failure_modes:
        dist = dist_by_fm.get(fm.id)
        if dist is None:
            continue
        dist_rows.append(
            {
                "equipment_tag": tags.get(fm.equipment_id, ""),
                "failure_mode_name": fm.name,
                "distribution_json": json.dumps(
                    dist.distribution.model_dump(mode="json"),
                    ensure_ascii=False,
                ),
                "source_type": str(dist.provenance.source_type),
                "confidence": str(dist.provenance.confidence),
                "source_reference": dist.provenance.source_reference,
                "generated_by": dist.provenance.generated_by,
            }
        )
    _write_rows(
        wb.create_sheet(s.SHEET_FAILURE_DISTRIBUTIONS),
        s.FAILURE_DISTRIBUTION_COLUMNS,
        dist_rows,
    )

    duration_by_task = {
        d.maintenance_task_id: d for d in content.maintenance_distributions
    }
    maint_rows: list[dict[str, Any]] = []
    for task in content.maintenance_tasks:
        duration = duration_by_task.get(task.id)
        row: dict[str, Any] = {
            "equipment_tag": tags.get(task.equipment_id, ""),
            "name": task.name,
            "task_type": str(task.task_type),
            "trigger": str(task.trigger),
            "failure_mode_name": (
                fm_names.get(task.failure_mode_id)
                if task.failure_mode_id
                else None
            ),
            "interval_value": (task.interval.value if task.interval else None),
            "interval_unit": (
                str(task.interval.unit) if task.interval else None
            ),
            "cost": task.cost,
        }
        if duration is not None:
            row["duration_distribution_json"] = json.dumps(
                duration.distribution.model_dump(mode="json"),
                ensure_ascii=False,
            )
            row["duration_source_type"] = str(duration.provenance.source_type)
            row["duration_confidence"] = str(duration.provenance.confidence)
        maint_rows.append(row)
    _write_rows(
        wb.create_sheet(s.SHEET_MAINTENANCE),
        s.MAINTENANCE_COLUMNS,
        maint_rows,
    )

    task_eq: dict[UUID, UUID] = {
        t.id: t.equipment_id for t in content.maintenance_tasks
    }
    effect_rows: list[dict[str, Any]] = []
    for effect in content.maintenance_effects:
        eq_id = task_eq.get(effect.maintenance_task_id)
        effect_rows.append(
            {
                "equipment_tag": tags.get(eq_id, "") if eq_id else "",
                "maintenance_name": maint_names.get(
                    effect.maintenance_task_id, ""
                ),
                "effect_type": str(effect.effect_type),
                "failure_mode_name": (
                    fm_names.get(effect.failure_mode_id)
                    if effect.failure_mode_id
                    else None
                ),
                "parameter": effect.parameter,
            }
        )
    _write_rows(
        wb.create_sheet(s.SHEET_MAINTENANCE_EFFECTS),
        s.MAINTENANCE_EFFECT_COLUMNS,
        effect_rows,
    )

    diag_rows = [
        {
            "equipment_tag": tags.get(diag.equipment_id, ""),
            "failure_mode_name": fm_names.get(diag.failure_mode_id, ""),
            "name": diag.name,
            "method": diag.method,
            "interval_value": diag.interval.value,
            "interval_unit": str(diag.interval.unit),
            "detection_probability": diag.detection_probability,
            "false_positive_probability": diag.false_positive_probability,
        }
        for diag in content.diagnostic_tasks
    ]
    _write_rows(
        wb.create_sheet(s.SHEET_DIAGNOSTICS),
        s.DIAGNOSTIC_COLUMNS,
        diag_rows,
    )

    resource_rows = [
        {
            "name": res.name,
            "resource_type": res.resource_type,
            "capacity": res.capacity,
            "cost_per_hour": res.cost_per_hour,
        }
        for res in content.resources
    ]
    _write_rows(
        wb.create_sheet(s.SHEET_RESOURCES),
        s.RESOURCE_COLUMNS,
        resource_rows,
    )

    spare_rows = [
        {
            "name": spare.name,
            "stock": spare.stock,
            "lead_time_value": (
                spare.lead_time.value if spare.lead_time else None
            ),
            "lead_time_unit": (
                str(spare.lead_time.unit) if spare.lead_time else None
            ),
            "unit_cost": spare.unit_cost,
        }
        for spare in content.spare_parts
    ]
    _write_rows(
        wb.create_sheet(s.SHEET_SPARE_PARTS),
        s.SPARE_PART_COLUMNS,
        spare_rows,
    )

    prod_rows: list[dict[str, Any]] = []
    for fn in content.production_functions:
        prod_rows.append(
            {
                "kind": "function",
                "product": fn.product,
                "nominal_rate_value": fn.nominal_rate.value,
                "mass_unit": str(fn.nominal_rate.mass_unit),
                "time_unit": str(fn.nominal_rate.time_unit),
            }
        )
    for impact in content.production_impacts:
        prod_rows.append(
            {
                "kind": "impact",
                "equipment_tag": tags.get(impact.equipment_id, ""),
                "failure_mode_name": (
                    fm_names.get(impact.failure_mode_id)
                    if impact.failure_mode_id
                    else None
                ),
                "loss_fraction": impact.loss_fraction,
            }
        )
    _write_rows(
        wb.create_sheet(s.SHEET_PRODUCTION),
        s.PRODUCTION_COLUMNS,
        prod_rows,
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
